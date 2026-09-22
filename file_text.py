"""
Извлечение текста из файлов лекций (PDF/DOCX/PPTX/TXT) — контекст для решалки
по лекциям (см. gemini_solver.py, database.get_subject_lecture_context, PLAN.md
Фаза 9). Работает поверх уже сохранённых в боте файлов: качает байты по
Telegram file_id (bot.get_file + bot.download_file) и парсит в обычный текст.

Не делает OCR картинок внутри PDF/PPTX — если у файла нет текстового слоя
(скан), extract_text вернёт пустую строку, и extract_and_save просто не
сохранит ничего для этого файла (сам файл при этом никуда не девается,
остаётся доступен как обычно через /files — теряется только его участие
в контексте лекций для решалки).
"""

import io
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".txt")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
    return "\n".join(parts)


def extract_text(data: bytes, file_name: str) -> str:
    """Возвращает текст файла, либо '' если формат не поддержан или парсинг
    не удался (битый файл, скан без текстового слоя и т.д.)."""
    name = (file_name or "").lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf(data)
        if name.endswith(".docx"):
            return _extract_docx(data)
        if name.endswith(".pptx"):
            return _extract_pptx(data)
        if name.endswith(".txt"):
            return data.decode("utf-8", errors="ignore")
    except Exception as e:
        logger.warning(f"Не смог распарсить {file_name}: {e}")
        return ""
    return ""


async def extract_and_save(bot, db_file_id: int, tg_file_id: str, file_name: str) -> bool:
    """Качает файл по Telegram file_id и сохраняет извлечённый текст в БД
    (привязано к db_file_id — id записи в таблице files). Возвращает True,
    если текст удалось извлечь и сохранить (файл войдёт в контекст лекций
    предмета для решалки), False — если формат не поддержан, файл пустой
    (скан) или скачивание/парсинг не удались.

    Ошибки намеренно не пробрасываются наружу: это вспомогательная фича
    поверх уже сохранённого файла (add_file к этому моменту уже отработал),
    её сбой не должен ломать сам аплоад/синк файлов."""
    name = (file_name or "").lower()
    if not name.endswith(SUPPORTED_EXTENSIONS):
        return False
    try:
        tg_file = await bot.get_file(tg_file_id)
        buf = await bot.download_file(tg_file.file_path)
        data = buf.read()
    except Exception as e:
        logger.warning(f"Не смог скачать файл для извлечения текста (file_id={db_file_id}): {e}")
        return False

    text = extract_text(data, file_name)
    if not text.strip():
        return False

    from database import save_file_text
    try:
        await save_file_text(db_file_id, text)
    except Exception as e:
        # Тот же контракт, что и выше: сбой не должен ронять аплоад/синк
        # (в /syncfiles исключение отсюда обрывало импорт всех оставшихся файлов).
        logger.warning(f"Не смог сохранить текст файла (file_id={db_file_id}): {e}")
        return False
    return True
