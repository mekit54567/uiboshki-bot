"""Поиск по смыслу (semantic_index + semantic_search): нарезка по страницам и
слайдам, путь «по словам» (FTS5/BM25) и «по смыслу» (векторы в sqlite-vec),
слияние RRF, разнообразие MMR, чат со ссылками «· слайд 12» и страница лекции.
Сеть не нужна: эмбеддинги Gemini подменены детерминированными векторами."""
import io

import pytest

import semantic_index as si
import semantic_search as ss

DIMS = 8
# «смыслы» для подменённых эмбеддингов: текст с любым из слов → своя ось
CONCEPTS = [("дисконт", "npv", "прибыль", "окупа"), ("регресс", "коэффициент"), ("баланс", "актив"), ("привет",)]


def fake_vec(text: str) -> list[float]:
    low = text.lower()
    v = [0.0] * DIMS
    for i, words in enumerate(CONCEPTS):
        if any(w in low for w in words):
            v[i] = 1.0
    v[DIMS - 1] = 0.15                      # общий «фон», чтобы не было нулевых векторов
    n = sum(x * x for x in v) ** 0.5
    return [x / n for x in v]


@pytest.fixture
def embed(monkeypatch):
    import config
    import gemini_solver
    calls = []

    async def fake(texts, task="RETRIEVAL_DOCUMENT", titles=None, dims=None, timeout=60):
        calls.append((task, len(texts)))
        return [fake_vec(t) for t in texts]

    monkeypatch.setattr(config, "GEMINI_EMBED_DIMS", DIMS)
    monkeypatch.setattr(gemini_solver, "embed", fake)
    ss._qcache.clear()
    return calls


def pptx_bytes(slides: list[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for text in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3)).text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pdf_bytes(pages: int) -> bytes:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument.new()
    for _ in range(pages):
        pdf.new_page(200, 280)
    buf = io.BytesIO()
    pdf.save(buf)
    return buf.getvalue()


def test_chunk_pages_keeps_places():
    long = "Дисконтирование денежных потоков. " * 80                    # ~2700 символов — режется
    pages = ["Тема 5. Инвестиции", "Спасибо за внимание", "", long, "Коротко про NPV"]
    chunks = si.chunk_pages(pages)
    assert chunks[0][:2] == (1, 2)                                        # два коротких слайда — вместе
    assert all(a == b == 4 for a, b, _ in chunks[1:-1]) and len(chunks) >= 4   # длинная — кусками, все со стр. 4
    assert all(len(t) <= si.CHUNK * 1.5 for _, _, t in chunks)
    assert chunks[-1][:2] == (5, 5)                                       # через пустую страницу не склеивает
    assert chunks[1][2][-40:] in "".join(t for _, _, t in chunks[1:3])     # перехлёст на стыке


def test_pages_of_pptx_and_docx():
    kind, pages = si.pages_of(pptx_bytes(["Первый слайд", "Второй слайд про NPV"]), "лк5.pptx")
    assert kind == "слайд" and pages == ["Первый слайд", "Второй слайд про NPV"]
    from docx import Document
    d = Document()
    for i in range(30):
        d.add_paragraph(f"Абзац {i} " + "текст " * 20)
    buf = io.BytesIO()
    d.save(buf)
    kind, parts = si.pages_of(buf.getvalue(), "метод.docx")
    assert kind == "часть" and len(parts) > 1 and all(len(p) <= 1700 for p in parts)
    assert si.pages_of(b"garbage", "x.pdf") == ("", [])


def test_rrf_and_mmr():
    fused = ss.rrf([1, 2, 3], [3, 4])
    assert max(fused, key=fused.get) == 3                                 # есть в обоих списках — выше
    a = {"id": 1, "file_id": 1, "score": 1.0, "text": "дисконтирование потоков", "vec": [1, 0]}
    dup = {"id": 2, "file_id": 2, "score": 0.99, "text": "дисконтирование потоков", "vec": [1, 0]}
    other = {"id": 3, "file_id": 3, "score": 0.7, "text": "регрессия", "vec": [0, 1]}
    assert [c["id"] for c in ss.mmr([a, dup, other], k=2)] == [1, 3]       # копия уступает другому смыслу
    same = [{"id": i, "file_id": 7, "score": 1 - i / 10, "text": f"кусок {i}"} for i in range(5)]
    assert len(ss.mmr(same, k=5, per_file=3)) == 3


@pytest.mark.asyncio
async def test_hybrid_search_finds_by_meaning(db, embed):
    f1 = await db.add_file("Лекция 5. Инвестиции", "Анализ данных", "tg1", "лк5.pptx", 1)
    f2 = await db.add_file("Лекция 2. Регрессия", "Анализ данных", "tg2", "лк2.pptx", 1)
    f3 = await db.add_file("Баланс предприятия", "Учётная деятельность", "tg3", "уч.pptx", 1)
    for fid in (f1, f2, f3):
        await db.save_file_text(fid, "текст")
    await si.index_file(f1, pptx_bytes(["Оценка проектов и сравнение вариантов. " * 7,
                                        "Дисконтирование и NPV: деньги дешевеют со временем. " * 6]), "лк5.pptx")
    await si.index_file(f2, pptx_bytes(["Линейная регрессия: коэффициент b. " * 8]), "лк2.pptx")
    await si.index_file(f3, pptx_bytes(["Актив и пассив баланса. " * 8]), "уч.pptx")
    assert await ss.ready()

    # до векторов — только по словам, но уже с нужным слайдом
    hits = await ss.search("что такое регрессия")
    assert hits[0]["file_id"] == f2 and hits[0]["lex"] == 1 and hits[0]["sem"] is None

    assert await si.embed_pending(pause=0) == 4 and embed[-1] == ("RETRIEVAL_DOCUMENT", 4)
    # ни одного общего слова с лекцией — находится по смыслу
    hits = await ss.search("почему нельзя просто сложить прибыль за все годы")
    assert hits[0]["file_id"] == f1 and hits[0]["lex"] is None and hits[0]["sem"] == 1
    assert ss.place(hits[0]) == "слайд 2" and hits[0]["sim"] > 0.9
    # предмет — фильтр: баланс ищем в «Анализе данных» — его там нет
    assert all(h["subject"] == "Анализ данных" for h in await ss.search("актив баланса", "Анализ данных"))
    # «привет» без предмета — лекции не тянем
    assert await ss.search("привет") == []

    ctx, sources = ss.build_context(hits)
    assert ctx.startswith("=== [1] Анализ данных: Лекция 5. Инвестиции · слайд 2 ===")
    assert sources[0] == {"n": 1, "id": f1, "title": "Лекция 5. Инвестиции", "page": 2, "kind": "слайд",
                          "label": "Лекция 5. Инвестиции · слайд 2"}
    from ai_solver import lecture_system_prompt
    assert "[1], [2]" in lecture_system_prompt("", ctx)                    # ИИ просят ссылаться на номера

    # файл удалили — его куски уходят из поиска при следующем проходе
    from database._conn import connect
    async with connect() as m:
        await m.execute("DELETE FROM file_text WHERE file_id=?", (f3,))
        await m.commit()
    st = await si.index_pending(None, max_batches=0)
    assert st["files"] == 2 and st["waiting"] == 0


@pytest.mark.asyncio
async def test_index_pending_downloads_for_pages(db, embed):
    fid = await db.add_file("Лекция 1", "Анализ данных", "tg-1", "лк1.pptx", 1)
    await db.save_file_text(fid, "Слайд один Слайд два")
    big = await db.add_file("Огромная", "Анализ данных", "tg-big", "big.pdf", 1)
    await db.save_file_text(big, "Текст огромного файла про регрессию")

    class Bot:
        async def get_file(self, tg_id):
            if tg_id == "tg-big":
                raise RuntimeError("file is too big")
            return type("F", (), {"file_path": tg_id})()

        async def download_file(self, path):
            return io.BytesIO(pptx_bytes(["Слайд один", "Слайд два про регрессию " * 12]))

    st = await si.index_pending(Bot(), max_batches=4)
    assert st["files"] == 2 and st["embedded"] == st["chunks"] > 0
    async with si.open_index() as idx:
        kinds = dict(await (await idx.execute("SELECT file_id, kind FROM chunks")).fetchall())
    assert kinds[fid] == "слайд" and kinds[big] == "часть"                # не скачался — по сохранённому тексту


@pytest.mark.asyncio
async def test_chat_cites_pages_and_page_view(db, embed, monkeypatch):
    from fastapi.testclient import TestClient
    import ai_solver
    import webapp.server as server
    from tests.test_webapp_auth import BOT_TOKEN, _make_init_data
    monkeypatch.setattr(server.deps, "BOT_TOKEN", BOT_TOKEN)
    fid = await db.add_file("Лекция 5. Инвестиции", "Анализ данных", "tg5", "лк5.pdf", 1)
    await db.save_file_text(fid, "текст")
    data = pdf_bytes(3)
    await si.index_file(fid, None, "лк5.pdf", "Дисконтирование и NPV " * 10)
    async with si.open_index() as idx:                                    # как будто это стр. 2 из 3
        await idx.execute("UPDATE chunks SET kind='стр.', page_from=2, page_to=2")
        await idx.execute("INSERT INTO chunks (file_id, page_from, page_to, kind, text) VALUES (?, 3, 3, 'стр.', 'Итоги')", (fid,))
        await idx.commit()
    await si.embed_pending(pause=0)
    seen = {}

    async def fake_chat(history, subject="", extra_system="", lectures=""):
        seen["lectures"] = lectures
        return {"content": "Будущие деньги дисконтируют [1].", "reasoning": ""}

    monkeypatch.setattr(ai_solver, "chat_with_reasoning", fake_chat)
    c, h = TestClient(server.app), {"X-Telegram-Init-Data": _make_init_data()}
    r = c.post("/api/chat", headers=h, json={"history": [{"role": "user", "content": "как учесть, что прибыль приходит позже"}]})
    assert r.status_code == 200, r.text
    res = r.json()
    assert seen["lectures"].startswith("=== [1]") and "[1]" in res["html"]
    assert res["sources"][0]["label"] == "Лекция 5. Инвестиции · стр. 2" and res["sources"][0]["page"] == 2

    p = c.get(f"/api/files/{fid}/page/2", headers=h).json()
    assert p["pages"] == 3 and p["kind"] == "стр." and "Дисконтирование" in p["text"] and p["image"]
    assert c.get(f"/api/files/{fid}/page/99", headers=h).json()["page"] == 3    # дальше последней — последняя

    class Bot:
        async def get_file(self, tg_id):
            return type("F", (), {"file_path": tg_id})()

        async def download_file(self, path):
            return io.BytesIO(data)

    monkeypatch.setattr(server.deps, "tg_bot", lambda: Bot())
    from urllib.parse import urlsplit
    u = urlsplit(p["image"])
    img = c.get(f"{u.path}?{u.query}")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg" and img.content[:2] == b"\xff\xd8"
    assert c.get(f"/pg/{fid}/2.jpg?exp=9999999999&sig=bad").status_code == 403


@pytest.mark.asyncio
async def test_local_model_own_vectors_and_switch_back(db, embed, monkeypatch):
    """Своя модель (переезд в Россию): свои векторы в своей таблице — Gemini
    не трогаем; заранее считаются фоном (EMBED_PRECOMPUTE), переключение
    EMBED_PROVIDER туда и обратно — без пересчёта; удалённый файл уходит
    из векторов всех моделей."""
    import local_embed
    calls = []

    async def fake_local(key, texts, query=False):
        calls.append((key, query, len(texts)))
        return [fake_vec(t) for t in texts]

    monkeypatch.setattr(local_embed, "embed", fake_local)
    monkeypatch.setitem(local_embed.MODELS, "rubert-mini-frida", {**local_embed.MODELS["rubert-mini-frida"], "dims": DIMS})
    monkeypatch.delenv("EMBED_PROVIDER", raising=False)
    monkeypatch.setenv("EMBED_PRECOMPUTE", "local:rubert-mini-frida")
    f1 = await db.add_file("Лекция 5. Инвестиции", "Анализ данных", "tg1", "лк5.pptx", 1)
    f2 = await db.add_file("Лекция 2. Регрессия", "Анализ данных", "tg2", "лк2.pptx", 1)
    for fid in (f1, f2):
        await db.save_file_text(fid, "текст")
    await si.index_file(f1, pptx_bytes(["Дисконтирование и NPV: деньги дешевеют со временем. " * 6]), "лк5.pptx")
    await si.index_file(f2, pptx_bytes(["Линейная регрессия: коэффициент b. " * 8]), "лк2.pptx")
    st = await si.index_pending(None, max_batches=4)            # Gemini + заранее своя
    assert st["provider"] == "gemini" and st["embedded"] == st["chunks"] == 2
    assert calls == [("rubert-mini-frida", False, 2)]

    monkeypatch.setenv("EMBED_PROVIDER", "rubert-mini-frida")    # переехали: ищет своя, Gemini не зовём
    embed.clear()
    hits = await ss.search("почему нельзя просто сложить прибыль за все годы")
    assert hits[0]["file_id"] == f1 and hits[0]["sem"] == 1 and embed == []
    assert calls[-1] == ("rubert-mini-frida", True, 1)
    assert (await si.stats())["embedded"] == 2

    monkeypatch.setenv("EMBED_PROVIDER", "gemini")               # и обратно — векторы Gemini на месте
    hits = await ss.search("почему нельзя просто сложить прибыль за все годы")
    assert hits[0]["file_id"] == f1 and hits[0]["sem"] == 1 and embed[-1] == ("RETRIEVAL_QUERY", 1)

    await si.index_file(f2, None, "лк2.pptx", text="")           # перенарезка — старые векторы всех моделей прочь
    async with si.open_index() as idx:
        left = (await (await idx.execute("SELECT COUNT(*) FROM emb_rubert_mini_frida")).fetchone())[0]
    assert left == 1
