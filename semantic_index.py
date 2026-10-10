"""
Индекс лекций для поиска по смыслу (поиск — semantic_search.py).

Лекция режется на куски с привязкой к месту: страница PDF, слайд PPTX,
часть DOCX/TXT. У каждого куска два «отпечатка»:
  • слова — FTS5 (BM25) прямо в SQLite;
  • смысл — вектор Gemini или своей модели на сервере (embedder.py,
    EMBED_PROVIDER) в виртуальной таблице vec0 расширения sqlite-vec:
    ближайшие по косинусу ищет сам SQLite. У каждой модели своя таблица.

Индекс — производные данные: лежит отдельным файлом рядом с базой
(<база>.search.db), в ночной бэкап не идёт (пересобирается из файлов), а
файлы и предметы берёт из основной базы через ATTACH — удалил староста файл,
и его куски уходят из поиска при следующем проходе.

Наполняется фоном (scheduler: index_pending): новые файлы режутся сразу при
загрузке (file_text.extract_and_save), старые — понемногу, а векторы
досчитываются пачками в пределах бесплатных лимитов Gemini; поиск работает и
до конца индексации — по словам, с нужной страницей.
"""

import io
import logging
import os
import re
import struct
from contextlib import asynccontextmanager

import aiosqlite

logger = logging.getLogger(__name__)

CHUNK = 1100          # символов в куске: абзац-два — смысл не размывается
MIN_PAGE = 220        # слайд короче — склеиваем с соседним
OVERLAP = 150         # перехлёст длинных страниц, чтобы мысль не рвалась на стыке
EMBED_BATCH = 64
KINDS = {".pdf": "стр.", ".pptx": "слайд"}    # у DOCX/TXT страниц нет — «часть»

VEC_OK: bool | None = None      # загрузилось ли sqlite-vec (None — ещё не пробовали)


def index_path() -> str:
    import database
    root, _ = os.path.splitext(database.DATABASE_PATH)
    return root + ".search.db"


# ── страницы и куски ──────────────────────────────────────────────────────────

def _norm(text: str) -> str:
    lines = [re.sub(r"[ \t ]+", " ", ln).strip() for ln in (text or "").splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(ln for ln in lines)).strip()


def text_parts(text: str, size: int = 1500) -> list[str]:
    """Текст без страниц (DOCX, TXT, старый сплошной текст) — на «части» по
    абзацам, около size символов."""
    parts, cur = [], ""
    for para in re.split(r"\n\s*\n|\n", text or ""):
        para = para.strip()
        if not para:
            continue
        if cur and len(cur) + len(para) + 1 > size:
            parts.append(cur)
            cur = para
        else:
            cur = f"{cur}\n{para}" if cur else para
    if cur:
        parts.append(cur)
    return parts


def pages_of(data: bytes, file_name: str) -> tuple[str, list[str]]:
    """(вид места, тексты по порядку): PDF — страницы, PPTX — слайды, прочее —
    части. Не разобралось — ("", [])."""
    from file_text import extract_text
    name = (file_name or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            return KINDS[".pdf"], [(p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages]
        if name.endswith(".pptx"):
            from pptx import Presentation
            slides = []
            for slide in Presentation(io.BytesIO(data)).slides:
                texts = [sh.text_frame.text.strip() for sh in slide.shapes
                         if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
                slides.append("\n".join(texts))
            return KINDS[".pptx"], slides
        text = extract_text(data, file_name)
        return ("часть", text_parts(text)) if text.strip() else ("", [])
    except Exception as e:
        logger.info(f"страницы {file_name}: {type(e).__name__}: {e}")
        return "", []


def _split_long(text: str) -> list[str]:
    """Длинная страница — на куски по предложениям, с перехлёстом OVERLAP."""
    sents = re.split(r"(?<=[.!?…;:])\s+|\n", text)
    out, cur = [], ""
    for s in sents:
        if not s.strip():
            continue
        if cur and len(cur) + len(s) + 1 > CHUNK:
            out.append(cur)
            tail = cur[-OVERLAP:]
            cur = tail[tail.find(" ") + 1:] + " " + s if " " in tail else s
        else:
            cur = f"{cur} {s}" if cur else s
        while len(cur) > CHUNK * 1.5:           # «предложение» без точек на полстраницы
            out.append(cur[:CHUNK])
            cur = cur[CHUNK - OVERLAP:]
    if cur.strip():
        out.append(cur)
    return out


def chunk_pages(pages: list[str]) -> list[tuple[int, int, str]]:
    """[(страница_с, страница_по, текст)]: короткие соседние страницы
    склеиваются (слайд «Спасибо за внимание» сам по себе бесполезен), длинные
    режутся. Кусок не перескакивает через пустую страницу и не длиннее CHUNK."""
    out: list[tuple[int, int, str]] = []
    buf, b_from, b_to = "", 0, 0

    def flush():
        nonlocal buf
        if buf.strip():
            out.append((b_from, b_to, buf))
        buf = ""

    for i, raw in enumerate(pages, 1):
        t = _norm(raw)
        if not t:
            flush()
            continue
        if len(t) > CHUNK:
            flush()
            out.extend((i, i, piece) for piece in _split_long(t))
            continue
        if buf and len(buf) + len(t) + 1 <= CHUNK and (len(buf) < MIN_PAGE or len(t) < MIN_PAGE):
            buf, b_to = buf + "\n" + t, i
        else:
            flush()
            buf, b_from, b_to = t, i, i
    flush()
    return out


# ── база индекса ──────────────────────────────────────────────────────────────

def _pack(vec: list[float]) -> bytes:
    return struct.pack(f"{len(vec)}f", *vec)


def unpack(blob: bytes) -> list[float]:
    return list(struct.unpack(f"{len(blob) // 4}f", blob))


async def _load_vec(db) -> bool:
    global VEC_OK
    if VEC_OK is False:
        return False
    try:
        import sqlite_vec
        await db.enable_load_extension(True)
        await db.load_extension(sqlite_vec.loadable_path())
        await db.enable_load_extension(False)
        VEC_OK = True
    except Exception as e:
        if VEC_OK is None:
            logger.warning(f"sqlite-vec не загрузился — поиск только по словам: {e}")
        VEC_OK = False
    return VEC_OK


async def _schema(db, dims: int):
    await db.executescript("""
        CREATE TABLE IF NOT EXISTS chunks (
            id        INTEGER PRIMARY KEY,
            file_id   INTEGER NOT NULL,
            page_from INTEGER NOT NULL,
            page_to   INTEGER NOT NULL,
            kind      TEXT NOT NULL,
            text      TEXT NOT NULL,
            embedded  INTEGER NOT NULL DEFAULT 0
        );
        CREATE INDEX IF NOT EXISTS chunks_file ON chunks(file_id);
        CREATE INDEX IF NOT EXISTS chunks_todo ON chunks(embedded);
        CREATE TABLE IF NOT EXISTS indexed_files (file_id INTEGER PRIMARY KEY, mode TEXT, at TEXT);
        CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
        CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
            text, content='chunks', content_rowid='id', tokenize='unicode61 remove_diacritics 2');
        CREATE TRIGGER IF NOT EXISTS chunks_ai AFTER INSERT ON chunks BEGIN
            INSERT INTO chunks_fts(rowid, text) VALUES (new.id, new.text);
        END;
        CREATE TRIGGER IF NOT EXISTS chunks_ad AFTER DELETE ON chunks BEGIN
            INSERT INTO chunks_fts(chunks_fts, rowid, text) VALUES ('delete', old.id, old.text);
        END;
    """)
    if VEC_OK:
        row = await (await db.execute("SELECT value FROM meta WHERE key='dims'")).fetchone()
        if row and int(row[0]) != dims:          # сменили длину вектора — старые векторы не годятся
            await db.execute("DROP TABLE IF EXISTS vec_chunks")
            await db.execute("UPDATE chunks SET embedded=0")
        await db.execute(f"CREATE VIRTUAL TABLE IF NOT EXISTS vec_chunks USING vec0("
                         f"embedding float[{int(dims)}] distance_metric=cosine)")
        await db.execute("INSERT OR REPLACE INTO meta (key, value) VALUES ('dims', ?)", (str(dims),))
        import embedder
        for p in {embedder.provider(), embedder.precompute()} - {embedder.GEMINI, None}:
            # своя модель — свои таблицы: векторы Gemini не трогаем, переключение обратно мгновенное
            await db.execute(f"CREATE VIRTUAL TABLE IF NOT EXISTS {embedder.vec_table(p)} USING vec0("
                             f"embedding float[{int(embedder.dims(p))}] distance_metric=cosine)")
            await db.execute(f"CREATE TABLE IF NOT EXISTS {embedder.done_table(p)} (id INTEGER PRIMARY KEY)")
    await db.commit()


async def _vector_tables(db) -> list[str]:
    """Все таблицы векторов (vec0) и отметок «посчитано» всех моделей."""
    rows = await (await db.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND "
        "((name LIKE 'vec\\_%' ESCAPE '\\' AND sql LIKE '%USING vec0%') OR name LIKE 'emb\\_%' ESCAPE '\\')")).fetchall()
    return [r[0] for r in rows]


async def has_vectors(db, p: str) -> bool:
    import embedder
    if not VEC_OK:
        return False
    done = embedder.done_table(p)
    sql = f"SELECT 1 FROM {done} LIMIT 1" if done else "SELECT 1 FROM chunks WHERE embedded=1 LIMIT 1"
    try:
        return bool(await (await db.execute(sql)).fetchone())
    except Exception:
        return False


@asynccontextmanager
async def open_index():
    """Подключение к индексу: sqlite-vec (если есть), схема, основная база
    как m — файлы и предметы берутся оттуда живыми."""
    import database
    from config import GEMINI_EMBED_DIMS
    db = await aiosqlite.connect(index_path())
    try:
        await _load_vec(db)
        await db.execute("ATTACH DATABASE ? AS m", (database.DATABASE_PATH,))
        await _schema(db, GEMINI_EMBED_DIMS)
        yield db
    finally:
        await db.close()


async def _replace_chunks(db, file_id: int, kind: str, chunks: list[tuple[int, int, str]], mode: str):
    old = [r[0] for r in await (await db.execute("SELECT id FROM chunks WHERE file_id=?", (file_id,))).fetchall()]
    if old and VEC_OK:
        marks = ",".join("?" * len(old))
        for table in await _vector_tables(db):      # номера кусков переиспользуются — старые векторы всех моделей прочь
            col = "id" if table.startswith("emb_") else "rowid"
            await db.execute(f"DELETE FROM {table} WHERE {col} IN ({marks})", old)
    await db.execute("DELETE FROM chunks WHERE file_id=?", (file_id,))
    await db.executemany("INSERT INTO chunks (file_id, page_from, page_to, kind, text) VALUES (?, ?, ?, ?, ?)",
                         [(file_id, a, b, kind, t) for a, b, t in chunks])
    await db.execute("INSERT OR REPLACE INTO indexed_files (file_id, mode, at) VALUES (?, ?, datetime('now'))",
                     (file_id, mode))


async def index_file(file_id: int, data: bytes | None = None, file_name: str = "", text: str = "") -> int:
    """Нарезать файл в индекс (векторы — потом, фоном). data — байты файла:
    тогда с номерами страниц и слайдов; без них — части сохранённого текста.
    → сколько кусков."""
    import asyncio
    kind, pages = ("", [])
    if data:
        kind, pages = await asyncio.to_thread(pages_of, data, file_name)
    mode = "pages"
    if not any(p.strip() for p in pages):
        if not text:
            from database._conn import connect
            async with connect() as m:
                row = await (await m.execute("SELECT content FROM file_text WHERE file_id=?", (file_id,))).fetchone()
            text = row[0] if row else ""
        kind, pages, mode = "часть", text_parts(text), "text"
    chunks = chunk_pages(pages)
    async with open_index() as db:
        await _replace_chunks(db, file_id, kind, chunks, mode)
        await db.commit()
    return len(chunks)


async def embed_pending(max_batches: int = 8, pause: float = 1.0, p: str | None = None) -> int:
    """Досчитать векторы кускам без них — пачками, пока не кончатся или Gemini
    не скажет «лимит» (тогда — в следующий проход). p — модель (embedder.py),
    по умолчанию та, что ищет сейчас. → сколько посчитано."""
    import asyncio
    import embedder
    p = p or embedder.provider()
    table, done_t = embedder.vec_table(p), embedder.done_table(p)
    todo = f"c.id NOT IN (SELECT id FROM {done_t})" if done_t else "c.embedded = 0"
    done = 0
    async with open_index() as db:
        if not VEC_OK:
            return 0
        for _ in range(max_batches):
            rows = await (await db.execute(
                "SELECT c.id, c.text, COALESCE(f.title, '') FROM chunks c LEFT JOIN m.files f ON f.id = c.file_id "
                f"WHERE {todo} ORDER BY c.id LIMIT ?", (EMBED_BATCH,))).fetchall()
            if not rows:
                break
            try:
                vecs = await embedder.embed([r[1] for r in rows], titles=[r[2] for r in rows], p=p)
            except embedder.EmbedError as e:
                logger.info(f"эмбеддинги: {e} — продолжу в следующий проход")
                break
            await db.executemany(f"INSERT OR REPLACE INTO {table} (rowid, embedding) VALUES (?, ?)",
                                 [(r[0], _pack(v)) for r, v in zip(rows, vecs)])
            if done_t:
                await db.executemany(f"INSERT OR IGNORE INTO {done_t} (id) VALUES (?)", [(r[0],) for r in rows])
            else:
                await db.executemany("UPDATE chunks SET embedded=1 WHERE id=?", [(r[0],) for r in rows])
            await db.commit()
            done += len(rows)
            if pause:
                await asyncio.sleep(pause)
    return done


async def index_pending(bot, max_files: int = 40, max_batches: int = 8) -> dict:
    """Фоновый проход (scheduler): убрать куски удалённых файлов, нарезать
    файлы с текстом, которых в индексе нет (PDF и PPTX — заново скачав, ради
    страниц), и досчитать векторы. → статистика для /status."""
    async with open_index() as db:
        gone = [r[0] for r in await (await db.execute(
            "SELECT file_id FROM indexed_files WHERE file_id NOT IN (SELECT file_id FROM m.file_text)")).fetchall()]
        for fid in gone:
            await _replace_chunks(db, fid, "", [], "gone")
            await db.execute("DELETE FROM indexed_files WHERE file_id=?", (fid,))
        todo = await (await db.execute(
            "SELECT f.id, f.file_id, COALESCE(f.file_name, '') FROM m.files f JOIN m.file_text t ON t.file_id = f.id "
            "WHERE f.id NOT IN (SELECT file_id FROM indexed_files) ORDER BY f.id LIMIT ?", (max_files,))).fetchall()
        await db.commit()
    for fid, tg_id, name in todo:
        data = None
        if bot is not None and name.lower().endswith(tuple(KINDS)):
            try:
                tg = await bot.get_file(tg_id)
                data = (await bot.download_file(tg.file_path)).read()
            except Exception as e:          # больше 20 МБ или файл недоступен — по сохранённому тексту
                logger.info(f"индекс: файл {fid} не скачался ({type(e).__name__}) — без страниц")
        try:
            await index_file(fid, data, name)
        except Exception as e:
            logger.warning(f"индекс: файл {fid}: {type(e).__name__}: {e}")
    embedded = await embed_pending(max_batches)
    import embedder
    if embedder.precompute():           # векторы другой модели заранее (до переезда)
        embedded += await embed_pending(max_batches, pause=0, p=embedder.precompute())
    st = await stats()
    try:
        import health
        await health.note("semantic", True, f"{st['chunks']} фрагментов из {st['files']} файлов, "
                                            f"с векторами {st['embedded']}")
    except Exception:
        pass
    return dict(st, new_files=len(todo), new_vectors=embedded)


async def stats() -> dict:
    async with open_index() as db:
        files = (await (await db.execute("SELECT COUNT(*) FROM indexed_files")).fetchone())[0]
        import embedder
        p = embedder.provider()
        if embedder.done_table(p):
            chunks = (await (await db.execute("SELECT COUNT(*) FROM chunks")).fetchone())[0]
            emb = (await (await db.execute(f"SELECT COUNT(*) FROM {embedder.done_table(p)}")).fetchone())[0] if VEC_OK else 0
        else:
            chunks, emb = await (await db.execute("SELECT COUNT(*), COALESCE(SUM(embedded), 0) FROM chunks")).fetchone()
        waiting = (await (await db.execute(
            "SELECT COUNT(*) FROM m.file_text WHERE file_id NOT IN (SELECT file_id FROM indexed_files)")).fetchone())[0]
    return {"files": files, "chunks": chunks, "embedded": emb, "waiting": waiting, "vectors": bool(VEC_OK),
            "provider": p}
